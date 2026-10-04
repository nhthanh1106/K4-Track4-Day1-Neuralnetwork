"""Persist run logs and populate the supplied experiment workbook template."""
from __future__ import annotations

import json
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _json_safe(value: Any):
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if hasattr(value, "item"):
        return value.item()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def save_result(result: dict, results_dir: str | Path = "../results") -> str:
    """Save config, histories, and summary without serializing model weights."""
    results_dir = Path(results_dir)
    results_dir.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    payload = {key: result[key] for key in ("cfg", "history", "summary")}
    path = results_dir / f"{exp_id}.json"
    path.write_text(json.dumps(_json_safe(payload), indent=2), encoding="utf-8")
    return str(path)


def load_results(results_dir: str | Path = "../results") -> list[dict]:
    results_dir = Path(results_dir)
    results = []
    for path in sorted(results_dir.glob("*.json")):
        results.append(json.loads(path.read_text(encoding="utf-8")))
    return sorted(results, key=lambda item: item["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    cfg = result["cfg"]
    summary = result["summary"]
    optimizer_names = {
        "sgd": "SGD",
        "sgd_momentum": "SGD+momentum",
        "adam": "Adam",
        "adamw": "AdamW",
    }
    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg.get("group", "other"),
        "description": cfg.get("description", cfg.get("hypothesis", "")),
        "loss": str(cfg.get("loss", "ce")).upper(),
        "optimizer": optimizer_names.get(cfg["optimizer"], cfg["optimizer"]),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": "-".join(str(width) for width in cfg.get("hidden", (256, 128))),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm"),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed"),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": "Y" if summary.get("diverged") else "N",
        "eval_acc": None if eval_scores is None else eval_scores.get("accuracy"),
        "eval_macro_f1": None if eval_scores is None else eval_scores.get("macro_f1"),
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str | Path, out_path: str | Path,
               seed_ids: list[str] | None = None,
               summary_notes: dict[str, str] | None = None) -> None:
    """Fill the template's input cells and retain its formulas, styling, and validation."""
    if len(rows) > 60:
        raise ValueError(f"template has 60 experiment rows; got {len(rows)}")
    template_path, out_path = Path(template_path), Path(out_path)
    workbook = load_workbook(template_path)
    ws = workbook["Experiments"]
    headers = [cell.value for cell in ws[1]]
    formula_headers = {
        "step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"
    }
    input_headers = [header for header in headers if header not in formula_headers]

    # Clear the template's example baseline row and any previously entered inputs.
    for row_index in range(2, 62):
        for column_index, header in enumerate(headers, start=1):
            if header not in formula_headers:
                ws.cell(row=row_index, column=column_index).value = None

    for row_index, record in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            if header in formula_headers:
                continue
            if header not in record:
                raise KeyError(f"Missing workbook field {header!r} for row {record.get('exp_id')}")
            value = record[header]
            ws.cell(row=row_index, column=column_index).value = value

    seeds = workbook["Seeds"]
    ids = seed_ids or []
    if len(ids) > 5:
        raise ValueError("the template has room for five baseline seed IDs")
    for offset in range(5):
        seeds.cell(row=2 + offset, column=1).value = ids[offset] if offset < len(ids) else None

    summary = workbook["Summary"]
    for row_index in range(2, 12):
        group = summary.cell(row=row_index, column=1).value
        if group and summary_notes and group in summary_notes:
            summary.cell(row=row_index, column=8).value = summary_notes[group]

    # Keep the supplied workbook structure while making long experiment descriptions,
    # notes, and metric columns readable when the completed sheet is opened.
    column_widths = {
        "A": 20, "B": 13, "C": 50, "D": 10, "E": 19, "F": 11, "G": 13,
        "H": 9, "I": 9, "J": 13, "K": 11, "L": 12, "M": 12, "N": 10,
        "O": 9, "P": 12, "Q": 14, "R": 12, "S": 14, "T": 14, "U": 11,
        "V": 14, "W": 14, "X": 13, "Y": 11, "Z": 12, "AA": 14,
        "AB": 31, "AC": 54,
    }
    for column, width in column_widths.items():
        ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 34
    numeric_formats = {
        "lr": "0.####", "weight_decay": "0.####", "dropout": "0.##",
        "step0_loss": "0.0000", "best_val_loss": "0.0000",
        "final_train_loss": "0.0000", "final_val_loss": "0.0000",
        "val_acc": "0.0000", "val_macro_f1": "0.0000",
        "time_per_epoch_s": "0.00", "peak_mem_MB": "0.0",
        "eval_acc": "0.0000", "eval_macro_f1": "0.0000",
    }
    for row_index in range(2, 2 + len(rows)):
        ws.row_dimensions[row_index].height = 38
        for column_index, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_index, column=column_index)
            if header in numeric_formats:
                cell.number_format = numeric_formats[header]
            if header in ("description", "notes"):
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment

    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["H"].width = 68
    summary.row_dimensions[1].height = 34
    for row_index in range(2, 12):
        summary.row_dimensions[row_index].height = 44
        note_cell = summary.cell(row=row_index, column=8)
        note_alignment = copy(note_cell.alignment)
        note_alignment.wrap_text = True
        note_alignment.vertical = "top"
        note_cell.alignment = note_alignment

    # Ensure Excel recalculates the retained template formulas when the file opens.
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
